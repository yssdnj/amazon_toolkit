from pathlib import Path
import math
import datetime
import sys

# ==============================================================================
# OS 自动路由：Windows 使用 xlwings（依赖 Excel），其他系统使用 openpyxl
# ==============================================================================
IS_WINDOWS = sys.platform == 'win32'

if IS_WINDOWS:
    import xlwings as xw
else:
    import re
    from copy import copy
    import pandas as pd
    from openpyxl import load_workbook

import pandas as pd  # fill_data / fill_session 两端都需要


# ══════════════════════════════════════════════════════════════════════════════
# 公共入口（run）
# ══════════════════════════════════════════════════════════════════════════════

def run():
    base_path = Path(__file__).resolve().parents[1]
    data_dir = base_path / "data" / "weekly_report_append"
    input_dir = data_dir / "input"
    output_dir = data_dir / "output"

    input_dir.mkdir(parents=True, exist_ok=True)
    output_dir.mkdir(parents=True, exist_ok=True)

    print(f"\n📂 输入文件目录: {input_dir}")
    print(f"📂 输出文件目录: {output_dir}")
    print(f"🖥  运行环境: {'Windows (xlwings)' if IS_WINDOWS else 'Linux/Mac (openpyxl)'}")

    week_number = input("请输入国家-周数（如 US-26W11）: ").strip()

    input_file = '周销售数据统计' + week_number + '.xlsx'
    output_file = '周销售数据统计' + week_number + '.xlsx'

    input_path = input_dir / input_file
    output_path = output_dir / output_file

    if not input_path.exists():
        print(f"❌ 文件不存在: {input_path}")
        return

    output_file = 'output_aggregated-' + week_number + '.xlsx'
    # ==============================================================================
    # 20260503修改：修复 source_file 路径 Bug
    # 旧代码：source_file = base_path / "data" / "weekly_report" / "output" / output_file
    #         dev 分支已将目录重命名为 weekly_data_clean，旧路径会导致文件找不到直接报错
    # 新代码：路径与 weekly_data_clean.py 的输出目录保持一致
    # ==============================================================================
    source_file = base_path / "data" / "weekly_data_clean" / "output" / output_file

    add_row(input_path, output_path)
    fill_data(source_file, output_path)
    fill_session(source_file, output_path)


# ══════════════════════════════════════════════════════════════════════════════
# Windows 实现（xlwings + Excel COM）
# ══════════════════════════════════════════════════════════════════════════════

def _win_add_row(file_path, save_path=None):

    sheets = ["SL","Toy", "ToyDH", "DL", "DSL", "MFL", "SFM", "SC","VHAN1","VHAN2","SL_访问量","Toy_访问量", "ToyDH_访问量", "DL_访问量", "DSL_访问量", "MFL_访问量", "SFM_访问量", "SC_访问量", "VHAN1_访问量", "VHAN2_访问量"]

    # ==============================================================================
    # 20260503修改：App 生命周期改用 try/finally 保护
    # 旧代码：app.quit() 直接写在函数末尾，若中途异常会导致 Excel 进程残留后台
    # 新代码：无论是否报错，finally 块都保证 wb.close() + app.quit() 被执行
    # ==============================================================================
    app = xw.App(visible=False)
    app.display_alerts = False
    app.screen_updating = False

    try:
        wb = app.books.open(file_path)

        today = datetime.date.today() - datetime.timedelta(days=7)

        monday = today - datetime.timedelta(days=today.weekday())
        sunday = monday + datetime.timedelta(days=6)

        year, week, _ = monday.isocalendar()

        year_week = f"{year}W{week:02d}"
        week_range = f"{monday.strftime('%m%d')}~{sunday.strftime('%m%d')}"

        for sheet_name in sheets:

            if sheet_name not in [s.name for s in wb.sheets]:
                print(f"sheet不存在: {sheet_name}")
                continue

            sht = wb.sheets[sheet_name]

            # 找最后一行
            last_row = sht.range("A" + str(sht.cells.last_cell.row)).end("up").row

            insert_row = last_row
            template_row = last_row - 1

            # 插入新行（倒数两行之间）
            sht.range(f"{insert_row}:{insert_row}").api.Insert()

            # 复制模板行
            sht.range(f"{template_row}:{template_row}").api.Copy()
            # 粘贴格式
            sht.range(f"{insert_row}:{insert_row}").api.PasteSpecial(-4122)
            # 粘贴公式
            sht.range(f"{insert_row}:{insert_row}").api.PasteSpecial(-4123)

            # # 删除常量单元格（只保留公式）
            # used_cols = sht.used_range.last_cell.column
            # for col in range(3, used_cols + 1):
            #     cell = sht.range(insert_row, col)
            #     if not cell.formula:
            #         cell.value = None

            # 填充前两个单元格
            sht.range(insert_row, 1).value = week_range
            sht.range(insert_row, 2).value = year_week

            print(f"完成 sheet: {sheet_name}")

        if save_path:
            wb.save(save_path)
        else:
            wb.save()

    finally:
        # 旧代码：wb.close() / app.quit() 直接在函数末尾调用，异常时不执行
        wb.close()
        app.quit()

    return "全部sheet处理完成"


def _win_fill_data(file1, file2):

    sheet_map = ["SL", "Toy", "ToyDH", "DL", "DSL", "MFL", "SFM", "SC", "VHAN1", "VHAN2"]

    summary_cols = [
        "销量", "订单量", "销售额", "促销销量", "促销订单量", "促销销售额",
        "促销折扣", "退款量", "退款金额", "展示", "点击",
        "广告订单量", "广告花费", "广告销售额", "CPC"
    ]

    # 读取附件1两个表
    df_summary = pd.read_excel(file1, sheet_name="标签汇总sht")
    df_product = pd.read_excel(file1, sheet_name="标签品名汇总sht")

    # 关键修复
    df_summary["listing标签"] = df_summary["listing标签"].astype(str).str.strip()
    df_product["listing标签"] = df_product["listing标签"].astype(str).str.strip()

    # ==============================================================================
    # 20260503修改：App 生命周期改用 try/finally 保护
    # 旧代码：wb.close() / app.quit() 直接写在函数末尾，异常时不会执行
    # 新代码：finally 块保证 Excel 进程一定被释放
    # ==============================================================================
    app = xw.App(visible=False)
    try:
        wb = app.books.open(file2)

        for tag in sheet_map:

            # ==============================================================================
            # 20260503修改：裸 except 改为 except Exception as e，避免掩盖真实错误
            # 旧代码：except:（捕获所有异常包括 KeyboardInterrupt，且不打印原因）
            # 新代码：except Exception as e，并把错误信息打印出来
            # ==============================================================================
            try:
                sht = wb.sheets[tag]
            except Exception as e:
                print(f"⚠ sheet不存在或打开失败: {tag}，原因: {e}")
                continue

            # 找写入行
            last_row = sht.used_range.last_cell.row
            write_row = last_row - 1

            headers = sht.range("A3").expand("right").value

            # 修复：整行写回会覆盖公式列，改为只对数据列逐个写入，公式列完全不动

            # 第一部分：标签汇总数据
            df_tag = df_summary[df_summary["listing标签"] == tag]

            if not df_tag.empty:
                row_data = df_tag.iloc[0]
                for col in summary_cols:
                    if col not in headers:
                        continue
                    col_idx = headers.index(col) + 1  # xlwings 列号从1开始
                    sht.cells(write_row, col_idx).value = row_data[col]

            # 第二部分：品名数据
            df_tag_product = df_product[df_product["listing标签"] == tag]

            if not df_tag_product.empty:
                for _, row in df_tag_product.iterrows():
                    product_name = str(row["品名"]).strip()
                    sales = row["销量"]
                    refund_qty = row["退款量"]

                    if product_name in headers:
                        col_idx = headers.index(product_name) + 1
                        sht.cells(write_row, col_idx).value = sales

                    refund_col = f"退款{product_name}"
                    if refund_col in headers:
                        col_idx = headers.index(refund_col) + 1
                        sht.cells(write_row, col_idx).value = refund_qty

            print(f"✓ 完成 sheet: {tag}")

        wb.save()

    finally:
        # 旧代码：wb.close() / app.quit() 直接写在函数末尾，异常时不执行
        wb.close()
        app.quit()

    print("✓ 全部完成")


def _win_fill_session(file1, file2):

    session_sheet_map = {
        "SL_访问量":    "SL",
        "Toy_访问量":   "Toy",
        "ToyDH_访问量": "ToyDH",
        "DL_访问量":    "DL",
        "DSL_访问量":   "DSL",
        "MFL_访问量":   "MFL",
        "SFM_访问量":   "SFM",
        "SC_访问量":    "SC",
        "VHAN1_访问量": "VHAN1",
        "VHAN2_访问量": "VHAN2"
    }

    session_cols = [
        "Sessions-Browser", "Sessions-Mobile", "Sessions-Total",
        "PV-Browser", "PV-Mobile", "PV-Total",
    ]

    # 读取 source 文件的标签汇总表
    df_summary = pd.read_excel(file1, sheet_name="标签汇总sht")
    df_summary["listing标签"] = df_summary["listing标签"].astype(str).str.strip()

    # 只保留 source 文件实际存在的列（兼容旧文件无此六列）
    available_cols = [c for c in session_cols if c in df_summary.columns]
    if not available_cols:
        print("ℹ️ source 文件中无 Sessions/PV 列，跳过 fill_session")
        return

    app = xw.App(visible=False)
    try:
        wb = app.books.open(file2)

        for sheet_name, tag in session_sheet_map.items():

            try:
                sht = wb.sheets[sheet_name]
            except Exception as e:
                print(f"⚠ sheet不存在或打开失败: {sheet_name}，原因: {e}")
                continue

            # 找写入行（与 fill_data 完全相同）
            last_row = sht.used_range.last_cell.row
            write_row = last_row - 1

            # 列头在第3行（与 fill_data 完全相同）
            headers = sht.range("A3").expand("right").value

            # 按 listing标签 找对应行，只写数据列，不读写整行（避免覆盖公式）
            df_tag = df_summary[df_summary["listing标签"] == tag]
            if not df_tag.empty:
                row_data = df_tag.iloc[0]
                for col in available_cols:
                    if col not in headers:
                        continue
                    val = row_data[col]
                    col_idx = headers.index(col) + 1  # xlwings 列号从1开始
                    if val is None or (isinstance(val, float) and math.isnan(val)):
                        sht.cells(write_row, col_idx).value = None
                    else:
                        sht.cells(write_row, col_idx).value = int(val)
            print(f"✓ 完成 sheet: {sheet_name}（{tag}）")

        wb.save()

    finally:
        wb.close()
        app.quit()

    print("✓ fill_session 全部完成")


# ══════════════════════════════════════════════════════════════════════════════
# Linux/Mac 实现（openpyxl，纯 Python，无需 Excel）
# ══════════════════════════════════════════════════════════════════════════════

def _adjust_formula(formula, row_offset):
    """调整公式中的相对行引用（对应 Excel 复制行时自动偏移行号的行为）
    例：=(C81-C80)/C80 向下复制1行 → =(C82-C81)/C81
    规则：相对行引用（无 $ 前缀）加 row_offset；绝对行引用（有 $ 前缀）不变
    """
    if not formula or not isinstance(formula, str) or not formula.startswith('='):
        return formula

    def replace_ref(match):
        col_abs = match.group(1)
        col     = match.group(2)
        row_abs = match.group(3)
        row     = match.group(4)
        if row_abs:
            return f"{col_abs}{col}{row_abs}{row}"
        else:
            return f"{col_abs}{col}{int(row) + row_offset}"

    return re.sub(r'(\$?)([A-Z]+)(\$?)(\d+)', replace_ref, formula)


def _get_last_row(ws):
    """找最后一个任意列有值的行号（扫描所有列，避免只看A列遗漏汇总行）"""
    for row in range(ws.max_row, 0, -1):
        for col in range(1, ws.max_column + 1):
            if ws.cell(row=row, column=col).value is not None:
                return row
    return 1


def _expand_table_ref(ws, inserted_row):
    """插入行后更新工作表中所有 Excel 表格（ListObject）的引用范围
    openpyxl 的 insert_rows() 不会自动更新 Table.ref，
    导致新行落在表格汇总行位置，需手动将表格末行扩展一行
    """
    for tbl in ws.tables.values():
        if ':' not in tbl.ref:
            continue
        start_ref, end_ref = tbl.ref.split(':')
        match = re.match(r'([A-Z]+)(\d+)$', end_ref)
        if not match:
            continue
        end_col, end_row = match.group(1), int(match.group(2))
        if end_row >= inserted_row:
            tbl.ref = f"{start_ref}:{end_col}{end_row + 1}"


def _update_summary_formulas(ws, summary_row, old_end_row, new_end_row):
    """更新汇总行的公式引用，处理两种情况：
    1. 数据范围末端扩展（SUM 类）：old_end_row → new_end_row
       例：=SUM(B3:B88) → =SUM(B3:B89)
    2. 汇总行自引用偏移（除法类，如 ACoS）：old_summary_row → summary_row
       例：插入前汇总行在第89行，=Z89/AA89 → 移到第90行后应变为 =Z90/AA90
    """
    old_summary_row = summary_row - 1  # 插入前汇总行所在位置

    def expand_ref(match):
        col_abs = match.group(1)
        col     = match.group(2)
        row_abs = match.group(3)
        row     = int(match.group(4))
        if row_abs:
            return match.group(0)
        if row == old_end_row:
            return f"{col_abs}{col}{new_end_row}"
        if row == old_summary_row:
            return f"{col_abs}{col}{summary_row}"
        return match.group(0)

    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=summary_row, column=col)
        if isinstance(cell.value, str) and cell.value.startswith('='):
            cell.value = re.sub(r'(\$?)([A-Z]+)(\$?)(\d+)', expand_ref, cell.value)


def _copy_row_style(ws, src_row, dst_row):
    """将 src_row 的格式和公式复制到 dst_row，并自动调整相对行引用
    对应 xlwings：api.Copy() + PasteSpecial(-4122 格式) + PasteSpecial(-4123 公式)
    """
    row_offset = dst_row - src_row
    for col in range(1, ws.max_column + 1):
        src = ws.cell(row=src_row, column=col)
        dst = ws.cell(row=dst_row, column=col)
        if isinstance(src.value, str) and src.value.startswith('='):
            dst.value = _adjust_formula(src.value, row_offset)
        else:
            dst.value = src.value
        if src.has_style:
            dst.font = copy(src.font)
            dst.fill = copy(src.fill)
            dst.border = copy(src.border)
            dst.alignment = copy(src.alignment)
            dst.number_format = src.number_format
            dst.protection = copy(src.protection)


def _get_headers(ws, header_row=3):
    """读取指定行表头，遇到空列停止（对应 xlwings range.expand('right').value）"""
    headers = []
    for col in range(1, ws.max_column + 1):
        val = ws.cell(row=header_row, column=col).value
        if val is None:
            break
        headers.append(val)
    return headers


def _linux_add_row(file_path, save_path=None):

    sheets = ["SL", "Toy", "ToyDH", "DL", "DSL", "MFL", "SFM", "SC", "VHAN1", "VHAN2",
              "SL_访问量", "Toy_访问量", "ToyDH_访问量", "DL_访问量",
              "DSL_访问量", "MFL_访问量", "SFM_访问量", "SC_访问量","VHAN1_访问量","VHAN2_访问量"]

    wb = load_workbook(file_path)

    today = datetime.date.today() - datetime.timedelta(days=7)
    monday = today - datetime.timedelta(days=today.weekday())
    sunday = monday + datetime.timedelta(days=6)
    year, week, _ = monday.isocalendar()
    year_week = f"{year}W{week:02d}"
    week_range = f"{monday.strftime('%m%d')}~{sunday.strftime('%m%d')}"

    for sheet_name in sheets:

        if sheet_name not in wb.sheetnames:
            print(f"sheet不存在: {sheet_name}")
            continue

        ws = wb[sheet_name]

        summary_row  = _get_last_row(ws)
        template_row = summary_row - 1
        insert_row   = summary_row

        ws.insert_rows(insert_row)
        _expand_table_ref(ws, insert_row)
        _copy_row_style(ws, template_row, insert_row)
        _update_summary_formulas(ws, summary_row + 1, template_row, insert_row)

        ws.cell(row=insert_row, column=1).value = week_range
        ws.cell(row=insert_row, column=2).value = year_week

        print(f"完成 sheet: {sheet_name}")

    if save_path:
        wb.save(save_path)
    else:
        wb.save(file_path)

    return "全部sheet处理完成"


def _linux_fill_data(file1, file2):

    sheet_map = ["SL", "Toy", "ToyDH", "DL", "DSL", "MFL", "SFM", "SC", "VHAN1", "VHAN2"]

    summary_cols = [
        "销量", "订单量", "销售额", "促销销量", "促销订单量", "促销销售额",
        "促销折扣", "退款量", "退款金额", "展示", "点击",
        "广告订单量", "广告花费", "广告销售额", "CPC"
    ]

    df_summary = pd.read_excel(file1, sheet_name="标签汇总sht")
    df_product = pd.read_excel(file1, sheet_name="标签品名汇总sht")

    df_summary["listing标签"] = df_summary["listing标签"].astype(str).str.strip()
    df_product["listing标签"] = df_product["listing标签"].astype(str).str.strip()

    wb = load_workbook(file2)

    for tag in sheet_map:

        if tag not in wb.sheetnames:
            print(f"⚠ sheet不存在或打开失败: {tag}")
            continue

        ws = wb[tag]
        last_row = ws.max_row
        write_row = last_row - 1
        headers = _get_headers(ws)

        df_tag = df_summary[df_summary["listing标签"] == tag]

        if not df_tag.empty:
            row_data = df_tag.iloc[0]
            for col in summary_cols:
                if col not in headers:
                    continue
                col_idx = headers.index(col) + 1
                ws.cell(row=write_row, column=col_idx).value = row_data[col]

        df_tag_product = df_product[df_product["listing标签"] == tag]

        if not df_tag_product.empty:
            for _, row in df_tag_product.iterrows():
                product_name = str(row["品名"]).strip()
                sales = row["销量"]
                refund_qty = row["退款量"]

                if product_name in headers:
                    col_idx = headers.index(product_name) + 1
                    ws.cell(row=write_row, column=col_idx).value = sales

                refund_col = f"退款{product_name}"
                if refund_col in headers:
                    col_idx = headers.index(refund_col) + 1
                    ws.cell(row=write_row, column=col_idx).value = refund_qty

        print(f"✓ 完成 sheet: {tag}")

    wb.save(file2)
    print("✓ 全部完成")


def _linux_fill_session(file1, file2):

    session_sheet_map = {
        "SL_访问量":    "SL",
        "Toy_访问量":   "Toy",
        "ToyDH_访问量": "ToyDH",
        "DL_访问量":    "DL",
        "DSL_访问量":   "DSL",
        "MFL_访问量":   "MFL",
        "SFM_访问量":   "SFM",
        "SC_访问量":    "SC",
        "VHAN1_访问量": "VHAN1",
        "VHAN2_访问量": "VHAN2"
    }

    session_cols = [
        "Sessions-Browser", "Sessions-Mobile", "Sessions-Total",
        "PV-Browser", "PV-Mobile", "PV-Total",
    ]

    df_summary = pd.read_excel(file1, sheet_name="标签汇总sht")
    df_summary["listing标签"] = df_summary["listing标签"].astype(str).str.strip()

    available_cols = [c for c in session_cols if c in df_summary.columns]
    if not available_cols:
        print("ℹ️ source 文件中无 Sessions/PV 列，跳过 fill_session")
        return

    wb = load_workbook(file2)

    for sheet_name, tag in session_sheet_map.items():

        if sheet_name not in wb.sheetnames:
            print(f"⚠ sheet不存在或打开失败: {sheet_name}")
            continue

        ws = wb[sheet_name]
        last_row = ws.max_row
        write_row = last_row - 1
        headers = _get_headers(ws)

        df_tag = df_summary[df_summary["listing标签"] == tag]
        if not df_tag.empty:
            row_data = df_tag.iloc[0]
            for col in available_cols:
                if col not in headers:
                    continue
                val = row_data[col]
                col_idx = headers.index(col) + 1
                if val is None or (isinstance(val, float) and math.isnan(val)):
                    ws.cell(row=write_row, column=col_idx).value = None
                else:
                    ws.cell(row=write_row, column=col_idx).value = int(val)

        print(f"✓ 完成 sheet: {sheet_name}（{tag}）")

    wb.save(file2)
    print("✓ fill_session 全部完成")


# ══════════════════════════════════════════════════════════════════════════════
# 统一对外接口，由 OS 自动路由到对应实现
# ══════════════════════════════════════════════════════════════════════════════

def add_row(file_path, save_path=None):
    if IS_WINDOWS:
        return _win_add_row(file_path, save_path)
    else:
        return _linux_add_row(file_path, save_path)


def fill_data(file1, file2):
    if IS_WINDOWS:
        return _win_fill_data(file1, file2)
    else:
        return _linux_fill_data(file1, file2)


def fill_session(file1, file2):
    if IS_WINDOWS:
        return _win_fill_session(file1, file2)
    else:
        return _linux_fill_session(file1, file2)
